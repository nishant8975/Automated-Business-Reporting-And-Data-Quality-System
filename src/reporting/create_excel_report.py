import os
import sys
import psycopg2
import pandas as pd
from dotenv import load_dotenv
import openpyxl
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.styles import Font, Alignment, PatternFill, numbers
from openpyxl.chart import BarChart, LineChart, Reference, PieChart
from openpyxl.worksheet.table import Table, TableStyleInfo
import shutil
from datetime import datetime

def get_conn():
    load_dotenv()
    return psycopg2.connect(
        dbname="automated_business_reporting",
        user=os.environ.get("PG_USER"),
        password=os.environ.get("PG_PASSWORD"),
        host=os.environ.get("PG_HOST"),
        port=os.environ.get("PG_PORT")
    )

def fetch_data(conn, view_name):
    query = f"SELECT * FROM analytics.{view_name}"
    return pd.read_sql(query, conn)

def apply_header_style(ws):
    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

def format_currency(ws, col_letters):
    for col in col_letters:
        for cell in ws[col]:
            if cell.row > 1:
                cell.number_format = '#,##0.00'

def format_number(ws, col_letters):
    for col in col_letters:
        for cell in ws[col]:
            if cell.row > 1:
                cell.number_format = '#,##0'

def main():
    conn = get_conn()
    
    # Load Data
    try:
        df_kpi = fetch_data(conn, 'v_order_kpis')
        df_monthly = fetch_data(conn, 'v_monthly_sales')
        df_category = fetch_data(conn, 'v_category_sales')
        df_payment = fetch_data(conn, 'v_payment_analysis')
        df_review = fetch_data(conn, 'v_review_analysis')
        df_status = fetch_data(conn, 'v_order_status')
        df_customer = fetch_data(conn, 'v_customer_analysis')
        df_dq = fetch_data(conn, 'v_data_quality_summary')
    except Exception as e:
        print(f"Error fetching data: {e}")
        sys.exit(1)
        
    conn.close()
    
    # Validations
    kpi_gmv = float(df_kpi['gmv'].iloc[0])
    kpi_orders = int(df_kpi['total_orders'].iloc[0])
    
    cat_gmv = float(df_category['gmv'].sum())
    month_gmv = float(df_monthly['gmv'].sum())
    month_orders = int(df_monthly['total_orders'].sum())
    
    errors = []
    if abs(kpi_gmv - cat_gmv) > 0.1: errors.append(f"Category GMV ({cat_gmv}) != KPI GMV ({kpi_gmv})")
    if abs(kpi_gmv - month_gmv) > 0.1: errors.append(f"Monthly GMV ({month_gmv}) != KPI GMV ({kpi_gmv})")
    if kpi_orders != month_orders: errors.append(f"Monthly Orders ({month_orders}) != KPI Orders ({kpi_orders})")
    
    if errors:
        print("VALIDATION FAILED:")
        for err in errors:
            print(" -", err)
        sys.exit(1)
        
    # Create Workbook
    wb = openpyxl.Workbook()
    
    # 1. Executive Summary
    ws_exec = wb.active
    ws_exec.title = "Executive Summary"
    
    ws_exec.append(["AUTOMATED BUSINESS REPORT"])
    ws_exec.append(["Business Performance & Data Quality"])
    ws_exec.append([])
    
    ws_exec.append(["Total Orders", "Customers", "GMV", "AOV"])
    ws_exec.append([
        df_kpi['total_orders'].iloc[0], 
        df_kpi['total_customers'].iloc[0], 
        df_kpi['gmv'].iloc[0], 
        df_kpi['aov'].iloc[0]
    ])
    
    ws_exec.append([])
    ws_exec.append(["Repeat Customers", "Items Sold", "Freight", "Avg Review"])
    ws_exec.append([
        df_kpi['repeat_customers'].iloc[0], 
        df_kpi['total_items'].iloc[0], 
        df_kpi['total_freight'].iloc[0], 
        df_kpi['average_review_score'].iloc[0]
    ])
    
    # Format Exec Summary
    ws_exec['A1'].font = Font(bold=True, size=16)
    ws_exec['A2'].font = Font(italic=True, size=12)
    
    for row in [4, 7]:
        for col in ['A', 'B', 'C', 'D']:
            ws_exec[f"{col}{row}"].font = Font(bold=True)
            ws_exec[f"{col}{row}"].fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
            
    ws_exec['C5'].number_format = '#,##0.00'
    ws_exec['D5'].number_format = '#,##0.00'
    ws_exec['C8'].number_format = '#,##0.00'
    ws_exec['A5'].number_format = '#,##0'
    ws_exec['B5'].number_format = '#,##0'
    ws_exec['A8'].number_format = '#,##0'
    ws_exec['B8'].number_format = '#,##0'
    ws_exec['D8'].number_format = '0.00'
    
    # 2. Monthly Sales
    ws_month = wb.create_sheet(title="Monthly Sales")
    for r in dataframe_to_rows(df_monthly, index=False, header=True):
        # Convert pandas timestamp to date string for excel
        if len(r) > 0 and isinstance(r[0], pd.Timestamp):
            r[0] = r[0].strftime('%Y-%m')
        ws_month.append(r)
    apply_header_style(ws_month)
    format_currency(ws_month, ['C', 'D', 'E', 'F'])
    format_number(ws_month, ['B'])
    
    # Chart: Monthly GMV
    chart_month = LineChart()
    chart_month.title = "Monthly GMV Trend"
    chart_month.y_axis.title = "GMV"
    chart_month.x_axis.title = "Month"
    data_month = Reference(ws_month, min_col=3, min_row=1, max_row=ws_month.max_row)
    cats_month = Reference(ws_month, min_col=1, min_row=2, max_row=ws_month.max_row)
    chart_month.add_data(data_month, titles_from_data=True)
    chart_month.set_categories(cats_month)
    ws_exec.add_chart(chart_month, "A11")
    
    # 3. Category Analysis
    ws_cat = wb.create_sheet(title="Category Analysis")
    for r in dataframe_to_rows(df_category, index=False, header=True):
        ws_cat.append(r)
    apply_header_style(ws_cat)
    format_currency(ws_cat, ['C', 'D', 'E'])
    format_number(ws_cat, ['B', 'F'])
    
    # Chart: Category GMV
    chart_cat = BarChart()
    chart_cat.type = "col"
    chart_cat.title = "Top 10 Categories by GMV"
    data_cat = Reference(ws_cat, min_col=3, min_row=1, max_row=11)
    cats_cat = Reference(ws_cat, min_col=1, min_row=2, max_row=11)
    chart_cat.add_data(data_cat, titles_from_data=True)
    chart_cat.set_categories(cats_cat)
    ws_exec.add_chart(chart_cat, "A26")
    
    # 4. Payment Analysis
    ws_pay = wb.create_sheet(title="Payment Analysis")
    for r in dataframe_to_rows(df_payment, index=False, header=True):
        ws_pay.append(r)
    apply_header_style(ws_pay)
    format_currency(ws_pay, ['D', 'E'])
    format_number(ws_pay, ['B', 'C'])
    
    # 5. Customer Analysis
    ws_cust = wb.create_sheet(title="Customer Analysis")
    
    # Create customer type summary
    cust_summary = df_customer.groupby('customer_type').agg(
        total_customers=('customer_unique_id', 'count'),
        total_gmv=('total_gmv', 'sum')
    ).reset_index()
    
    ws_cust.append(['Customer Type', 'Total Customers', 'Total GMV'])
    for r in dataframe_to_rows(cust_summary, index=False, header=False):
        ws_cust.append(r)
        
    ws_cust.append([])
    ws_cust.append(['customer_unique_id', 'total_orders', 'total_gmv', 'first_order_date', 'last_order_date', 'customer_type'])
    for r in dataframe_to_rows(df_customer.head(500), index=False, header=False): # limit to 500 to avoid bloat
        ws_cust.append(r)
    
    format_currency(ws_cust, ['C'])
    format_number(ws_cust, ['B'])
    
    # 6. Review Analysis
    ws_rev = wb.create_sheet(title="Review Analysis")
    for r in dataframe_to_rows(df_review, index=False, header=True):
        ws_rev.append(r)
    apply_header_style(ws_rev)
    format_number(ws_rev, ['B'])
    
    # 7. Order Status
    ws_status = wb.create_sheet(title="Order Status")
    for r in dataframe_to_rows(df_status, index=False, header=True):
        ws_status.append(r)
    apply_header_style(ws_status)
    format_number(ws_status, ['B'])
    
    for cell in ws_status['C']:
        if cell.row > 1:
            cell.number_format = '0.00"%"'
            
    # Chart: Order Status
    chart_status = PieChart()
    chart_status.title = "Order Status Distribution"
    data_status = Reference(ws_status, min_col=2, min_row=1, max_row=ws_status.max_row)
    cats_status = Reference(ws_status, min_col=1, min_row=2, max_row=ws_status.max_row)
    chart_status.add_data(data_status, titles_from_data=True)
    chart_status.set_categories(cats_status)
    ws_exec.add_chart(chart_status, "J11")
            
    # 8. Data Quality
    ws_dq = wb.create_sheet(title="Data Quality")
    ws_dq.append(["Data quality exceptions were handled during the Python/PostgreSQL cleaning layer. The Excel report presents the validated results."])
    ws_dq.append([])
    
    for r in dataframe_to_rows(df_dq.melt(var_name='Metric', value_name='Value'), index=False, header=True):
        ws_dq.append(r)
    apply_header_style(ws_dq)
    
    # 9. Source Data
    ws_src = wb.create_sheet(title="Source Data")
    sources = [
        ["Metric / Report Section", "Source View", "Grain", "Purpose"],
        ["Executive KPI", "analytics.v_order_kpis", "Grand Total", "Overall business KPIs"],
        ["Monthly Sales", "analytics.v_monthly_sales", "Order Month", "Monthly performance"],
        ["Category Analysis", "analytics.v_category_sales", "Product Category", "Category performance"],
        ["Payment Analysis", "analytics.v_payment_analysis", "Payment Type", "Payment behavior"],
        ["Customer Analysis", "analytics.v_customer_analysis", "Customer", "Customer behavior"],
        ["Review Analysis", "analytics.v_review_analysis", "Review Score", "Customer feedback"],
        ["Order Status", "analytics.v_order_status", "Order Status", "Fulfillment distribution"],
        ["Data Quality", "analytics.v_data_quality_summary", "Grand Total", "Quality monitoring"]
    ]
    for row in sources:
        ws_src.append(row)
    apply_header_style(ws_src)
    
    # Save Workbook
    out_dir = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(__file__))), 'reports', 'excel')
    os.makedirs(out_dir, exist_ok=True)
    out_path = os.path.join(out_dir, 'automated_business_report.xlsx')
    wb.save(out_path)
    
    # Archive Logic
    archive_dir = os.path.join(out_dir, 'archive')
    os.makedirs(archive_dir, exist_ok=True)
    timestamp_str = datetime.now().strftime('%Y-%m-%d_%H-%M-%S')
    archive_path = os.path.join(archive_dir, f'automated_business_report_{timestamp_str}.xlsx')
    
    # Collision safety
    counter = 1
    while os.path.exists(archive_path):
        archive_path = os.path.join(archive_dir, f'automated_business_report_{timestamp_str}_{counter}.xlsx')
        counter += 1
        
    shutil.copy2(out_path, archive_path)
    
    print(f"\nSUCCESS: Excel Report generated at {out_path}")
    print(f"ARCHIVED: {archive_path}")
    print("Worksheets Created:", wb.sheetnames)
    print("Reconciliation Passed: True")

if __name__ == "__main__":
    main()
